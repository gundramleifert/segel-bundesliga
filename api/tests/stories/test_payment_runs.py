"""Paying through the bank (Stories F-7, F-4): the club's account, one SEPA file over the
approved claims, and the bank's answer — booked, or returned.

The file is pain.001.001.09 and is validated against its XSD when it is built, so these
tests read it back as XML and check what a bank would book: debtor, creditors, amounts.
"""

from __future__ import annotations

import xml.etree.ElementTree as ET
from uuid import uuid4

from app.models.auth import Grant, Relation
from tests.stories.test_login_and_roles import _problem_code
from tests.stories.test_reimbursements import IBAN, login, submitted, world

CLUB_IBAN = "DE02120300000000202051"
SECOND_IBAN = "DE02500105170137075030"
NS = {"p": "urn:iso:std:iso:20022:tech:xsd:pain.001.001.09"}


async def setup(client, caplog):
    """A club with a treasurer and a saved account, and two members with approved claims
    — one of them with two."""
    w = await world()
    treasurer = await login(
        client,
        caplog,
        f"t-{uuid4().hex[:6]}@example.org",
        Grant(relation=Relation.TREASURER, club_id=w.club),
    )
    saved = await client.put(
        f"/api/clubs/{w.club}/bank-account",
        headers=treasurer,
        json={"holder": "Host Sailing Club", "iban": CLUB_IBAN},
    )
    assert saved.status_code == 200, saved.text

    anna = await login(
        client,
        caplog,
        f"anna-{uuid4().hex[:6]}@example.org",
        Grant(relation=Relation.MEMBER, club_id=w.club),
    )
    ben = await login(
        client,
        caplog,
        f"ben-{uuid4().hex[:6]}@example.org",
        Grant(relation=Relation.MEMBER, club_id=w.club),
    )
    claims = [
        await submitted(client, anna, club=w.club),
        await submitted(client, anna, club=w.club),
        await submitted(client, ben, club=w.club, iban=SECOND_IBAN),
    ]
    for claim in claims:
        response = await client.post(
            f"/api/claims/{claim['id']}/approve", headers=treasurer, json={"items": []}
        )
        assert response.status_code == 200, response.text
    return w, treasurer, anna, claims


class TestClubBankAccount:
    """Story F-7: the account the club pays from — its manager's or treasurer's."""

    async def test_only_the_clubs_money_people_read_and_write_it(self, client, caplog):
        w = await world()
        manager = await login(
            client,
            caplog,
            f"m-{uuid4().hex[:6]}@example.org",
            Grant(relation=Relation.MANAGER, club_id=w.club),
        )
        member = await login(
            client,
            caplog,
            f"x-{uuid4().hex[:6]}@example.org",
            Grant(relation=Relation.MEMBER, club_id=w.club),
        )
        url = f"/api/clubs/{w.club}/bank-account"

        assert (await client.get(url, headers=manager)).json() is None
        saved = await client.put(
            url, headers=manager, json={"holder": "Host SC", "iban": "de02 1203 0000 0000 2020 51"}
        )
        assert saved.json()["iban"] == CLUB_IBAN

        for method in ("get", "put"):
            kwargs = {"json": {"holder": "x", "iban": CLUB_IBAN}} if method == "put" else {}
            response = await getattr(client, method)(url, headers=member, **kwargs)
            assert response.status_code == 403
        wrong = await client.put(
            url, headers=manager, json={"holder": "x", "iban": IBAN[:-1] + "1"}
        )
        assert _problem_code(wrong) == "iban-invalid"


class TestExportingARun:
    """Story F-7: one run, one payment per payee, one file the bank takes."""

    async def test_a_run_pays_every_open_claim_once_per_payee(self, client, caplog):
        w, treasurer, _, claims = await setup(client, caplog)

        run = await client.post(
            f"/api/clubs/{w.club}/payment-runs",
            headers=treasurer,
            json={"execution_date": "2026-10-09"},
        )
        assert run.status_code == 201, run.text
        body = run.json()
        # Anna's two claims are one transfer; Ben's is another.
        assert sorted(p["amount_cents"] for p in body["payments"]) == [1250, 2500]
        assert body["total_cents"] == 3750 and body["open"] is True

        for claim in claims:
            state = (await client.get(f"/api/claims/{claim['id']}", headers=treasurer)).json()
            assert state["payment_state"] == "in_payment"
        # Nothing is left to pay, so nothing is listed to pay.
        pending = (await client.get("/api/claims/pending", headers=treasurer)).json()
        assert [r for r in pending if r["club_id"] == w.club] == []
        empty = await client.post(f"/api/clubs/{w.club}/payment-runs", headers=treasurer, json={})
        assert _problem_code(empty) == "payment-run-empty"

    async def test_the_file_is_pain_001_09_with_the_clubs_account(self, client, caplog):
        w, treasurer, _, claims = await setup(client, caplog)
        run = (
            await client.post(f"/api/clubs/{w.club}/payment-runs", headers=treasurer, json={})
        ).json()

        response = await client.get(f"/api/payment-runs/{run['id']}/sepa", headers=treasurer)
        assert response.status_code == 200
        assert "attachment" in response.headers["content-disposition"]
        root = ET.fromstring(response.content)
        assert root.find(".//p:DbtrAcct/p:Id/p:IBAN", NS).text == CLUB_IBAN
        creditors = {e.text for e in root.findall(".//p:CdtrAcct/p:Id/p:IBAN", NS)}
        assert creditors == {IBAN, SECOND_IBAN}
        assert root.find(".//p:GrpHdr/p:CtrlSum", NS).text == "37.50"
        remittance = " ".join(e.text for e in root.findall(".//p:RmtInf/p:Ustrd", NS))
        for claim in claims:
            assert f"C-{claim['id']}" in remittance

        # The same message id again: a bank refuses the second upload of one run.
        again = await client.get(f"/api/payment-runs/{run['id']}/sepa", headers=treasurer)
        message_id = ".//p:GrpHdr/p:MsgId"
        assert (
            ET.fromstring(again.content).find(message_id, NS).text == root.find(message_id, NS).text
        )

    async def test_a_run_is_refused_without_the_clubs_account(self, client, caplog):
        w = await world()
        treasurer = await login(
            client,
            caplog,
            f"t-{uuid4().hex[:6]}@example.org",
            Grant(relation=Relation.TREASURER, club_id=w.club),
        )
        response = await client.post(
            f"/api/clubs/{w.club}/payment-runs", headers=treasurer, json={}
        )
        assert _problem_code(response) == "club-bank-account-missing"

    async def test_nobody_exports_their_own_claim(self, client, caplog):
        """Four eyes hold for the file too: a treasurer who is also a member cannot put
        their own approved claim into a run."""
        w, _, _, _ = await setup(client, caplog)
        both = await login(
            client,
            caplog,
            f"tm-{uuid4().hex[:6]}@example.org",
            Grant(relation=Relation.TREASURER, club_id=w.club),
            Grant(relation=Relation.MEMBER, club_id=w.club),
        )
        manager = await login(
            client,
            caplog,
            f"m-{uuid4().hex[:6]}@example.org",
            Grant(relation=Relation.MANAGER, club_id=w.club),
        )
        own = await submitted(client, both, club=w.club)
        await client.post(f"/api/claims/{own['id']}/approve", headers=manager, json={"items": []})

        response = await client.post(
            f"/api/clubs/{w.club}/payment-runs", headers=both, json={"claim_ids": [own["id"]]}
        )
        assert response.status_code == 403
        assert _problem_code(response) == "claim-own-decision"
        # Left to choose, the run simply leaves it out.
        run = await client.post(f"/api/clubs/{w.club}/payment-runs", headers=both, json={})
        claim_ids = {c for p in run.json()["payments"] for c in p["claim_ids"]}
        assert own["id"] not in claim_ids

    async def test_only_the_clubs_money_people_make_a_run(self, client, caplog):
        w, _, anna, claims = await setup(client, caplog)
        response = await client.post(
            f"/api/clubs/{w.club}/payment-runs", headers=anna, json={"claim_ids": [claims[2]["id"]]}
        )
        assert response.status_code == 403

    async def test_a_claim_not_approved_or_of_another_payer_is_refused(self, client, caplog):
        w, treasurer, anna, _ = await setup(client, caplog)
        fresh = await submitted(client, anna, club=w.club)
        response = await client.post(
            f"/api/clubs/{w.club}/payment-runs",
            headers=treasurer,
            json={"claim_ids": [fresh["id"]]},
        )
        assert _problem_code(response) == "claim-not-payable"


class TestSettlingARun:
    """Stories F-7, F-4: the bank's answer — booked settles the run, a return fails one
    payment and opens its claims again."""

    async def test_booked_pays_and_a_return_opens_the_claim_again(self, client, caplog):
        w, treasurer, _, claims = await setup(client, caplog)
        run = (
            await client.post(f"/api/clubs/{w.club}/payment-runs", headers=treasurer, json={})
        ).json()
        listed = (await client.get("/api/payment-runs", headers=treasurer)).json()
        assert run["id"] in [r["id"] for r in listed]

        bens = next(p for p in run["payments"] if p["amount_cents"] == 1250)
        failed = await client.post(
            f"/api/payments/{bens['id']}/fail",
            headers=treasurer,
            json={"reason": "AC04 account closed"},
        )
        assert failed.status_code == 200, failed.text

        booked = await client.post(f"/api/payment-runs/{run['id']}/settle", headers=treasurer)
        assert booked.status_code == 200, booked.text
        assert booked.json()["open"] is False

        states = {
            c["id"]: (await client.get(f"/api/claims/{c['id']}", headers=treasurer)).json()[
                "payment_state"
            ]
            for c in claims
        }
        assert states == {
            claims[0]["id"]: "paid",
            claims[1]["id"]: "paid",
            claims[2]["id"]: "unpaid",
        }
        pending = (await client.get("/api/claims/pending", headers=treasurer)).json()
        assert [(r["id"], r["action"]) for r in pending] == [(claims[2]["id"], "pay")]
        assert run["id"] not in [
            r["id"] for r in (await client.get("/api/payment-runs", headers=treasurer)).json()
        ]


class TestSpreadsheets:
    """Story F-7: the run and the year's claims as Excel files, for the club's books."""

    async def test_the_run_as_a_spreadsheet(self, client, caplog):
        w, treasurer, _, claims = await setup(client, caplog)
        run = (
            await client.post(f"/api/clubs/{w.club}/payment-runs", headers=treasurer, json={})
        ).json()

        response = await client.get(f"/api/payment-runs/{run['id']}/xlsx", headers=treasurer)
        assert response.status_code == 200
        rows = _rows(response.content)
        header, body = rows[0], rows[1:]
        assert "IBAN" in header
        assert {row[header.index("Claim")] for row in body} == {f"C-{c['id']}" for c in claims}
        assert sum(row[header.index("Amount (EUR)")] for row in body) == 37.5

    async def test_the_years_claims_as_a_spreadsheet(self, client, caplog):
        w, treasurer, anna, claims = await setup(client, caplog)
        await client.post("/api/claims", headers=anna, json={"title": "draft", "club_id": w.club})

        url = f"/api/clubs/{w.club}/claims.xlsx"
        rows = _rows((await client.get(url, headers=treasurer)).content)
        header, body = rows[0], rows[1:]
        # One row per item; each claim here has one. The draft is not in the books.
        assert sorted(row[header.index("Claim")] for row in body) == sorted(
            f"C-{c['id']}" for c in claims
        )
        assert (await client.get(url, headers=anna)).status_code == 403
        last_year = _rows((await client.get(url, headers=treasurer, params={"year": 2000})).content)
        assert last_year[1:] == []


def _rows(content: bytes) -> list[list]:
    from io import BytesIO

    from openpyxl import load_workbook

    sheet = load_workbook(BytesIO(content), read_only=True).active
    return [list(row) for row in sheet.iter_rows(values_only=True)]
